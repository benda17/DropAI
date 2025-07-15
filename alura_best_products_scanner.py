import os
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from consts import BEST_PRODUCTS_URL, EXCEL_FILE, SHEET_NAME, HEADERS

def safe_int(text: str):
    try:
        return int(text.replace(",", "").strip())
    except:
        return None

def safe_number(text: str):
    """
    Parse a number that might be integer or float, stripping commas.
    Returns int, float, or None on failure.
    """
    s = text.replace(",", "").strip()
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return None

def scrape_best_sellers(url: str):
    resp = requests.get(url, headers=HEADERS)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")

    items = []
    for div in soup.find_all("div", {"role": "listitem"}):
        a = div.find("a", href=True)
        if not a:
            continue

        # rank
        rank_div = a.select_one(".bs-table-text-rank")
        rank = safe_int(rank_div.get_text()) if rank_div else None

        # image URL (from <img> tag)
        img_tag = a.find("img")
        image_url = None
        if img_tag:
            image_url = img_tag.get("src") or img_tag.get("data-src")

        # collect all bs-table-text divs:
        # 0=headline, 1=total sales, 2=monthly sales,
        # 3=currency symbol, 4=revenue amount, 5=category
        txts = a.select(".bs-table-text")
        if len(txts) < 6:
            continue

        headline           = txts[0].get_text(strip=True)
        total_sales_raw    = txts[1].get_text(strip=True)
        monthly_sales_raw  = txts[2].get_text(strip=True)
        # revenue parts, in case needed later:
        revenue_symbol     = txts[3].get_text(strip=True)
        revenue_amount_raw = txts[4].get_text(strip=True)
        category           = txts[5].get_text(strip=True)

        total_sales    = safe_int(total_sales_raw) or total_sales_raw
        monthly_sales  = safe_int(monthly_sales_raw) or monthly_sales_raw
        revenue_amount = safe_number(revenue_amount_raw) or revenue_amount_raw
        revenue = {"symbol": revenue_symbol, "amount": revenue_amount}

        items.append({
            "rank":          rank,
            "headline":      headline,
            "link":          a["href"].strip(),
            "image":         image_url,
            "sales":         total_sales,
            "monthly_sales": monthly_sales,
            "category":      category
            # 'revenue' is available if you need it later
        })

    return items

def append_to_excel(items):
    # create file & header row if needed
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append([
            "Rank", "Headline", "Link", "image_url",
            "Sales", "Monthly Sales", "Category"
        ])
    else:
        wb = load_workbook(EXCEL_FILE)
        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
        else:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append([
                "Rank", "Headline", "Link", "image_url",
                "Sales", "Monthly Sales", "Category"
            ])

    # collect existing links to avoid duplicates
    existing_links = {
        row[2].value  # link is in the 3rd column (index 2)
        for row in ws.iter_rows(min_row=2, max_col=3)
    }

    # append new items
    new_count = 0
    for it in items:
        if it["link"] in existing_links:
            continue
        ws.append([
            it["rank"],
            it["headline"],
            it["link"],
            it["image"],
            it["sales"],
            it["monthly_sales"],
            it["category"]
        ])
        new_count += 1

    wb.save(EXCEL_FILE)
    print(f"Appended {new_count} new items to '{EXCEL_FILE}'.")

if __name__ == "__main__":
    items = scrape_best_sellers(BEST_PRODUCTS_URL)
    append_to_excel(items)
  
